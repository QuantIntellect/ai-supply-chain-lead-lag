"""Build SIE.DE_kpis.csv (FY2022Q1-FY2026Q3) from Siemens' own Excel results
files (segment figures + income statement, as published each quarter) and
regex digests of the earnings-release text (book-to-bill, backlog, outlook,
data-center commentary).  Figures are as reported in that quarter's release
(not restated later)."""
import os, sys, re, json, csv, warnings
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from irlib import *
import openpyxl
warnings.filterwarnings("ignore")

T = "SIE.DE"
RAW = os.path.join(RAW_IR, T)
TXT = os.path.join(TEXT_IR, T)
QEND = {1: "12-31", 2: "03-31", 3: "06-30", 4: "09-30"}

ev = {}
for r in csv.DictReader(open(os.path.join(PROC_IR, f"{T}_events.csv"))):
    if r["event_type"] in ("Q1", "Q2", "Q3", "Q4"):
        ev[r["fiscal_period"]] = r["event_date"]
man = {}
for r in csv.DictReader(open(os.path.join(PROC_IR, f"{T}_manifest.csv"))):
    man[(r["fiscal_period"], r["doc_type"])] = r["url"]


def num(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def seg_table(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sh = [s for s in wb.sheetnames if "egment" in s][0]
    rows = [list(r) for r in wb[sh].iter_rows(values_only=True)]
    hdr = None
    for i, r in enumerate(rows):
        cells = [str(c) if c is not None else "" for c in r]
        if any(c.strip() == "Orders" for c in cells) and any(c.strip() == "Revenue" for c in cells):
            hdr = cells
            break
    oi = hdr.index([c for c in hdr if c.strip() == "Orders"][0])
    ri = hdr.index([c for c in hdr if c.strip() == "Revenue"][0])
    pi = [j for j, c in enumerate(hdr) if c.strip().startswith("Profit") and "margin" not in c][0]
    mi = [j for j, c in enumerate(hdr) if c.strip().startswith("Profit margin")][0]
    out = {}
    for r in rows:
        name = str(r[0] or "").strip()
        if not name or num(r[oi]) is None:
            continue
        key = ("DI" if name.startswith("Digital Industr") else "SI" if name.startswith("Smart Infra") else
               "MO" if name.startswith("Mobility") else "SHS" if name.startswith("Siemens Healthin") else
               "IB" if name.startswith("Industrial Busin") else "GROUP" if name.startswith("Siemens (cont") else None)
        if not key:
            continue
        out[key] = {"orders": num(r[oi]), "orders_prior": num(r[oi + 1]), "orders_act": num(r[oi + 2]),
                    "orders_comp": num(r[oi + 3]), "rev": num(r[ri]), "rev_prior": num(r[ri + 1]),
                    "rev_act": num(r[ri + 2]), "rev_comp": num(r[ri + 3]), "profit": num(r[pi]),
                    "margin": num(r[mi])}
    inc = [s for s in wb.sheetnames if "ncome" in s][0]
    gp = rv = None
    for r in wb[inc].iter_rows(max_row=20, values_only=True):
        name = str(r[0] or "").strip()
        vals = [num(v) for v in r[1:] if num(v) is not None]
        if name == "Revenue" and vals:
            rv = vals[0]
        if name.startswith("Gross profit") and vals:
            gp = vals[0]
    out["gross_margin"] = round(100 * gp / rv, 1) if gp and rv else None
    return out


def pct(x):
    return None if x is None else round(100 * x, 1)


def text(p, dt):
    fp = os.path.join(TXT, f"{p}_{dt}.txt")
    return re.sub(r"\s+", " ", open(fp, encoding="utf-8", errors="ignore").read()) if os.path.exists(fp) else ""


# Short quotes from the quarter's analyst presentation (text grepped from the
# extracted presentation; wording kept close to source).
PRES_DC = {
    "2023Q2": "[pres] SI: Electrification up 34%; large orders from data center, semis & power dist. customers; record SI backlog ~EUR16bn",
    "2023Q3": "[pres] 'Winning in data centers through a tailored vertical approach'; 'Accelerating use of AI applications drives data center investment'",
    "2023Q4": "[pres] 'Accelerated adoption of generative AI drives data center investment'; SI several large orders from data center customers; record SI backlog EUR16.5bn",
    "2024Q1": "[pres] SI: several large orders from data center customers; U.S. orders supported by data center wins; record SI backlog EUR17bn",
    "2024Q2": "[pres] 'AI demand is a significant inflection for accelerated data center investment'; 'Unprecedented AI growth drives massive build out of data centers globally'; long-lead electrical supply chains under pressure",
    "2024Q3": "[pres] SI: 'Massive AI driven momentum from data center customers'; all-time high SI backlog EUR18.6bn",
    "2024Q4": "[pres] SI: Electrification orders +31%, driven by large orders from data center and energy customers",
    "2025Q1": "[pres] Compass Datacenters multi-year capacity agreement (standardized MV power skids); Electrification orders +10% on data center, energy, industry orders",
    "2025Q2": "[pres] SI: large data center orders below record prior-year quarter; Electrical Products orders -16%",
    "2025Q3": "[pres] U.S.: lower large data center orders in Electrical Products",
    "2025Q4": "[pres] U.S. orders +18%, strongly up in Electrification & EP on data center wins; CMD: Data Centers & AI factories target-market CAGR 11%",
    "2026Q1": "[pres] Data Center revenue growth ~35%; nVent liquid-cooling reference architecture for NVIDIA AI data centers",
    "2026Q2": "[pres] SI orders +35% to record, data centers with triple-digit growth; Data Center H1 revenue growth >45%; US$165m U.S. investment for AI infrastructure manufacturing",
    "2026Q3": "[pres] SI orders +42% to record, data centers up triple-digit; DC orders 9M FY26 up triple-digit to ~EUR6bn, DC revenue 9M >50% to EUR3.1bn; serving 9 of top 10 data center providers",
}
rows = []
for fy in range(2022, 2027):
    for q in range(1, 5):
        p = f"{fy}Q{q}"
        if p > "2026Q3":
            continue
        xl = os.path.join(RAW, f"{p}_financial_results_xlsx.xlsx")
        if not os.path.exists(xl):
            print("no xlsx", p)
            continue
        s = seg_table(xl)
        g, di, si, ib = s.get("GROUP", {}), s.get("DI", {}), s.get("SI", {}), s.get("IB", {})
        er = text(p, "earnings_release")
        pr = text(p, "press_release")
        m = re.search(r"Book-to-bill ratio (\d\.\d\d)", er)
        btb = float(m.group(1)) if m else (round(g["orders"] / g["rev"], 2) if g.get("rev") else None)
        btb_src = "earnings release table" if m else "computed orders/revenue"
        m = re.search(r"Order backlog \(in billions of €\) (\d+(?:\.\d)?)", er)
        backlog_bn = float(m.group(1)) if m else None
        # outlook (group): take the text from the last 'Outlook' heading
        oi = er.rfind("Outlook")
        ol = er[oi:oi + 4000] if oi >= 0 else ""
        lo = hi = None
        grp_sent = ""
        sents = re.split(r"(?<=[a-z0-9%)])\.\s", ol)
        for sent in sents:
            if re.search(r"Siemens [Gg]roup", sent) and re.search(r"comparable revenue|revenue growth", sent):
                if not grp_sent:
                    grp_sent = sent.strip()
                m = re.search(r"(-?\d+(?:\.\d)?)\s*%\s*to\s*(-?\d+(?:\.\d)?)\s*%", sent)
                if m:
                    lo, hi = float(m.group(1)), float(m.group(2))
                    break
        eps = None
        for sent in sents:
            if "EPS pre PPA" in sent or "purchase price allocation" in sent:
                rngs = re.findall(r"€\s?(\d+\.\d\d)\s*(?:to|–|-)\s*€?\s?(\d+\.\d\d)", sent)
                if rngs:
                    j = len(rngs) - 1 if re.search(r"from a range of .{0,40}to a range of", sent) else 0
                    eps = (float(rngs[j][0]), float(rngs[j][1]))
                    break
        m = re.search(r"Smart Infrastructure[^.]{0,120}?comparable revenue growth[^.]{0,60}?(\d+(?:\.\d)?)\s*%\s*to\s*(\d+(?:\.\d)?)\s*%", ol)
        si_g = (float(m.group(1)), float(m.group(2))) if m else None
        gt = []
        if lo is not None:
            gt.append(f"Group comparable revenue growth {lo:g}% to {hi:g}%")
            if "lower end" in grp_sent and "Siemens Group" in grp_sent:
                gt.append("group revenue growth expected at lower end of range")
        elif grp_sent:
            gt.append("Group: " + grp_sent[:160])
        if eps:
            gt.append(f"EPS pre PPA EUR {eps[0]:.2f}-{eps[1]:.2f}")
        if si_g:
            gt.append(f"SI comp rev growth {si_g[0]:g}-{si_g[1]:g}%")
        # data-center commentary (first two distinct snippets from release / press release)
        dc = []
        for src in (er, pr):
            for mm in re.finditer(r"data cent", src, re.I):
                a = src.rfind("•", 0, mm.start())
                a = a + 1 if a >= 0 and mm.start() - a < 260 else max(0, mm.start() - 160)
                b = src.find("•", mm.end())
                b = b if 0 <= b - mm.end() < 200 else mm.end() + 120
                snip = src[a:b].strip()
                if not any(snip[:60] in d for d in dc):
                    dc.append(snip[:300])
                if len(dc) >= 2:
                    break
            if len(dc) >= 2:
                break
        extra = {
            "orders_comp_pct": pct(g.get("orders_comp")), "revenue_comp_pct": pct(g.get("rev_comp")),
            "ib_profit_eur_m": ib.get("profit"), "ib_profit_margin_pct": pct(ib.get("margin")),
            "DI_orders": di.get("orders"), "DI_orders_yoy_pct": pct(di.get("orders_act")), "DI_orders_comp_pct": pct(di.get("orders_comp")),
            "DI_revenue": di.get("rev"), "DI_revenue_comp_pct": pct(di.get("rev_comp")), "DI_margin_pct": pct(di.get("margin")),
            "SI_orders": si.get("orders"), "SI_orders_yoy_pct": pct(si.get("orders_act")), "SI_orders_comp_pct": pct(si.get("orders_comp")),
            "SI_revenue": si.get("rev"), "SI_revenue_comp_pct": pct(si.get("rev_comp")), "SI_margin_pct": pct(si.get("margin")),
            "SI_book_to_bill": round(si["orders"] / si["rev"], 2) if si.get("rev") else None,
            "backlog_eur_bn": backlog_bn, "book_to_bill_source": btb_src,
            "gross_margin_source": "gross profit / revenue, consolidated income statement (quarter)",
            "guidance_rev_unit": "comparable revenue growth % (fiscal year)",
            "guidance_eps_pre_ppa": list(eps) if eps else None,
            "SI_guidance_comp_growth_pct": list(si_g) if si_g else None,
            "basis": "as reported in the quarter's own release",
        }
        rows.append({
            "fiscal_period": p, "period_end": f"{fy - 1 if q == 1 else fy}-{QEND[q]}",
            "release_date": ev.get(p, ""), "currency": "EUR", "unit": "EUR m",
            "revenue": g.get("rev"), "revenue_yoy_pct": pct(g.get("rev_act")),
            "orders": g.get("orders"), "orders_yoy_pct": pct(g.get("orders_act")),
            "book_to_bill": btb, "backlog": backlog_bn * 1000 if backlog_bn else None,
            "gross_margin_pct": s.get("gross_margin"), "guidance_text": "; ".join(gt),
            "guidance_rev_low": lo, "guidance_rev_high": hi, "ai_dc_notes": " | ".join(dc + ([PRES_DC[p]] if p in PRES_DC else [])),
            "extra_json": extra,
            "source_doc": man.get((p, "financial_results_xlsx"), "") + " ; " + man.get((p, "earnings_release"), ""),
        })
        print(p, g.get("orders"), g.get("rev"), btb, backlog_bn, s.get("gross_margin"), lo, hi, eps, si.get("orders"), len(dc))
print(write_kpis(T, rows), len(rows))
